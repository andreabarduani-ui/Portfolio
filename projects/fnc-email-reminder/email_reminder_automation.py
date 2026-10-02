"""Prepare and send individual FNC 3 course-completion reminder emails."""

import argparse
import csv
import html
import os
import re
import smtplib
import ssl
from dataclasses import dataclass
from email.message import EmailMessage
from pathlib import Path


SUPPORT_EMAIL = "a.barduani@accademiainformatica.com"
SUBJECT = "Progetto FNC 3 – contenuti formativi da completare"
VADEMECUM_URL = "https://emooc.it/docs/Vademecum_discenti_FNC.pdf"
COMPANY_NAME = "Accademia Informatica s.r.l."
COMPANY_WEBSITE = "http://www.accademiainformatica.com"
ITALIAN_CONFIDENTIALITY = (
    "Questo messaggio di posta elettronica contiene informazioni di carattere "
    "confidenziale rivolte esclusivamente al destinatario sopra indicato. La diffusione, "
    "distribuzione e/o copiatura del documento trasmesso da parte di qualsiasi soggetto "
    "diverso dal destinatario è proibita, sia ai sensi dell’art. 616 c.p., che ai sensi "
    "del Regolamento UE 2016/679. Nel caso aveste ricevuto questo messaggio di posta "
    "elettronica per errore, siete pregati di segnalarlo immediatamente a mezzo e-mail: "
    "gdpr@accademiainformatica.com e di distruggere quanto ricevuto (compresi i files "
    "allegati), senza farne copia. Ringraziamo anticipatamente per la vostra preziosa "
    "collaborazione."
)
ENGLISH_CONFIDENTIALITY = (
    "This e-mail message contains confidential information intended recipient only "
    "indicated above. Dissemination, distribution and / or copying of the document "
    "forwarded by any person other than the recipient is prohibited, pursuant to "
    "art. 616 of the Criminal Code, that pursuant to GDPR 2016/679. If you have received "
    "this email in error, please notify us immediately by e-mail: "
    "gdpr@accademiainformatica.com and destroy what has been received (including "
    "attached files), without making a copy. thank you in advance for your precious "
    "collaboration."
)
EMAIL_PATTERN = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")
ALIASES = {
    "name": ("nome", "name", "allievo", "nome allievo"),
    "email": ("email", "e-mail", "mail", "indirizzo email"),
    "courses": ("corsi", "courses", "corsi da finire", "corsi da completare", "video"),
}


@dataclass(frozen=True)
class Participant:
    name: str
    email: str
    courses: tuple[str, ...]


def _read_rows(path):
    if path.suffix.lower() == ".csv":
        with path.open(encoding="utf-8-sig", newline="") as source:
            sample = source.read(4096)
            source.seek(0)
            try:
                dialect = csv.Sniffer().sniff(sample, delimiters=";,\t")
            except csv.Error:
                dialect = csv.excel
                dialect.delimiter = ";"
            return list(csv.reader(source, dialect))

    if path.suffix.lower() in {".xlsx", ".xlsm"}:
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise ValueError("Per leggere Excel, installare openpyxl: pip install openpyxl") from exc
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            return list(workbook.active.iter_rows(values_only=True))
        finally:
            workbook.close()

    raise ValueError("Formato file non supportato: usare un file .csv, .xlsx o .xlsm.")


def _column_index(headers, field, override=None):
    if override:
        wanted = override.strip().casefold()
        for index, header in enumerate(headers):
            if str(header or "").strip().casefold() == wanted:
                return index
        raise ValueError(f"Colonna '{override}' non trovata.")

    normalized = [str(header or "").strip().casefold() for header in headers]
    for alias in ALIASES[field]:
        if alias in normalized:
            return normalized.index(alias)
    expected = ", ".join(ALIASES[field])
    raise ValueError(f"Colonna {field} non trovata. Intestazioni riconosciute: {expected}.")


def _cell(row, index):
    return str(row[index] if index < len(row) and row[index] is not None else "").strip()


def _parse_courses(value):
    return tuple(
        dict.fromkeys(
            course.strip().lstrip("•").strip()
            for course in re.split(r"\r?\n|[|;]", value)
            if course.strip().lstrip("•").strip()
        )
    )


def load_participants(path, name_column=None, email_column=None, courses_column=None):
    rows = _read_rows(Path(path))
    if not rows:
        raise ValueError("Il file è vuoto.")

    headers = rows[0]
    indexes = {
        "name": _column_index(headers, "name", name_column),
        "email": _column_index(headers, "email", email_column),
        "courses": _column_index(headers, "courses", courses_column),
    }
    participants = []
    seen_emails = set()
    errors = []

    for row_number, row in enumerate(rows[1:], start=2):
        if not any(str(value or "").strip() for value in row):
            continue
        name = _cell(row, indexes["name"]).replace("\r", " ").replace("\n", " ")
        email = _cell(row, indexes["email"])
        courses = _parse_courses(_cell(row, indexes["courses"]))
        normalized_email = email.casefold()
        issues = []

        if not name:
            issues.append("nome mancante")
        if not EMAIL_PATTERN.fullmatch(email):
            issues.append("email non valida")
        elif normalized_email in seen_emails:
            issues.append("email duplicata")
        else:
            seen_emails.add(normalized_email)
        if not courses:
            issues.append("nessun corso indicato")

        if issues:
            errors.append(f"Riga {row_number}: {', '.join(issues)}")
        else:
            participants.append(Participant(name, email, courses))

    if errors:
        raise ValueError("Correggere i dati prima di procedere:\n- " + "\n- ".join(errors))
    if not participants:
        raise ValueError("Nessun partecipante valido trovato.")
    return participants


def build_message(participant, sender):
    message = EmailMessage()
    message["Subject"] = SUBJECT
    message["From"] = sender
    message["To"] = participant.email

    course_list_text = "\n".join(f"- {course}" for course in participant.courses)
    course_list_html = "".join(f"<li>{html.escape(course)}</li>" for course in participant.courses)
    message.set_content(
        f"""Gentile {participant.name},

nell'ambito del Progetto FNC 3, la invitiamo a completare quanto prima i contenuti formativi che risultano ancora da avviare o in fase di fruizione.

Di seguito trova l'elenco dei video da completare:

{course_list_text}

Le ricordiamo alcune indicazioni importanti per garantire la corretta tracciabilità delle attività formative e la validità delle ore maturate:
- fruire dei contenuti con una connessione internet stabile e, ove possibile, tramite collegamento via cavo;
- non utilizzare strumenti, estensioni del browser o software che alterino la velocità di riproduzione dei video;
- svolgere le attività formative esclusivamente durante il regolare orario di lavoro.

Per ulteriori dettagli, può consultare il Vademecum dei Discenti:
{VADEMECUM_URL}

Per necessità o richieste di supporto, può contattarci all'indirizzo {SUPPORT_EMAIL}.

Ringraziandola per la collaborazione, porgiamo cordiali saluti.

Team Accademia Informatica
{SUPPORT_EMAIL}

{COMPANY_NAME}

Viale F. Tommaso Marinetti, 221 - 00143 Roma
Tel: 351 7669357 - 06 39746618 | Fax: 06 97749271
{COMPANY_WEBSITE}

{ITALIAN_CONFIDENTIALITY}

*************

{ENGLISH_CONFIDENTIALITY}
"""
    )
    message.add_alternative(
        f"""<html><body>
<p>Gentile {html.escape(participant.name)},</p>
<p>nell'ambito del <strong>Progetto FNC 3</strong>, la invitiamo a completare quanto prima i contenuti formativi che risultano ancora da avviare o in fase di fruizione.</p>
<p>Di seguito trova l'elenco dei video da completare:</p>
<ul>{course_list_html}</ul>
<p>Le ricordiamo alcune indicazioni importanti per garantire la corretta tracciabilità delle attività formative e la validità delle ore maturate:</p>
<ul>
<li>fruire dei contenuti con una connessione internet stabile e, ove possibile, tramite collegamento via cavo;</li>
<li>non utilizzare strumenti, estensioni del browser o software che alterino la velocità di riproduzione dei video;</li>
<li>svolgere le attività formative esclusivamente durante il regolare orario di lavoro.</li>
</ul>
<p>Per ulteriori dettagli, può consultare il <strong>Vademecum dei Discenti</strong>:<br>
<a href="{VADEMECUM_URL}">{VADEMECUM_URL}</a></p>
<p>Per necessità o richieste di supporto, può contattarci all'indirizzo
<a href="mailto:{SUPPORT_EMAIL}">{SUPPORT_EMAIL}</a>.</p>
<hr>
<p>Ringraziandola per la collaborazione, porgiamo cordiali saluti.</p>
<p>Team Accademia Informatica<br>
<a href="mailto:{SUPPORT_EMAIL}">{SUPPORT_EMAIL}</a></p>
<p><strong>{COMPANY_NAME}</strong><br>
Viale F. Tommaso Marinetti, 221 - 00143 Roma<br>
Tel: 351 7669357 - 06 39746618 | Fax: 06 97749271<br>
<a href="{COMPANY_WEBSITE}">www.accademiainformatica.com</a></p>
<p><small>{ITALIAN_CONFIDENTIALITY.replace("gdpr@accademiainformatica.com", '<a href="mailto:gdpr@accademiainformatica.com">gdpr@accademiainformatica.com</a>')}</small></p>
<p>*************</p>
<p><small>{ENGLISH_CONFIDENTIALITY.replace("gdpr@accademiainformatica.com", '<a href="mailto:gdpr@accademiainformatica.com">gdpr@accademiainformatica.com</a>')}</small></p>
</body></html>""",
        subtype="html",
    )
    return message


def _smtp_settings():
    host = os.environ.get("SMTP_HOST", "").strip()
    username = os.environ.get("SMTP_USER", "").strip()
    password = os.environ.get("SMTP_PASSWORD", "")
    sender = os.environ.get("EMAIL_FROM", username).strip()
    if not host:
        raise ValueError("Impostare SMTP_HOST nell'ambiente prima di inviare.")
    if bool(username) != bool(password):
        raise ValueError("SMTP_USER e SMTP_PASSWORD devono essere impostati insieme.")
    if not EMAIL_PATTERN.fullmatch(sender):
        raise ValueError("Impostare EMAIL_FROM (o SMTP_USER) con un indirizzo email valido.")
    return host, int(os.environ.get("SMTP_PORT", "587")), username, password, sender


def send_messages(participants, settings):
    host, port, username, password, sender = settings
    context = ssl.create_default_context()
    sent = 0
    failed = []
    with smtplib.SMTP(host, port, timeout=30) as server:
        server.ehlo()
        server.starttls(context=context)
        server.ehlo()
        if username:
            server.login(username, password)
        for participant in participants:
            try:
                refused = server.send_message(build_message(participant, sender))
                if refused:
                    reason = str(refused.get(participant.email, "Destinatario rifiutato"))
                    failed.append((participant.email, reason))
                    print(f"Invio non riuscito per {participant.email}: {reason}")
                else:
                    sent += 1
                    print(f"Inviata a {participant.name} <{participant.email}>")
            except (OSError, smtplib.SMTPException) as exc:
                failed.append((participant.email, str(exc)))
                print(f"Invio non riuscito per {participant.email}: {exc}")
    return sent, failed


def main(argv=None):
    parser = argparse.ArgumentParser(
        description="Prepara promemoria email individuali per i partecipanti al Progetto FNC 3."
    )
    parser.add_argument("file", type=Path, help="File Excel (.xlsx/.xlsm) o CSV con dati partecipanti.")
    parser.add_argument("--name-column", help="Intestazione della colonna con il nome.")
    parser.add_argument("--email-column", help="Intestazione della colonna con l'email.")
    parser.add_argument("--courses-column", help="Intestazione della colonna con i corsi da finire.")
    parser.add_argument(
        "--send",
        action="store_true",
        help="Invia i messaggi; senza questa opzione mostra solo l'anteprima.",
    )
    args = parser.parse_args(argv)

    try:
        participants = load_participants(
            args.file, args.name_column, args.email_column, args.courses_column
        )
        if not args.send:
            print(f"Anteprima: {len(participants)} email pronte. Nessuna email è stata inviata.")
            for person in participants:
                print(f"\nA: {person.name} <{person.email}>\nOggetto: {SUBJECT}")
                for course in person.courses:
                    print(f"  - {course}")
            print("\nVerificare i dati; usare --send per avviare l'invio.")
            return 0

        settings = _smtp_settings()
        print(f"Stai per inviare {len(participants)} email da {settings[4]}.")
        confirmation = input("Per confermare, digitare INVIA: ").strip()
        if confirmation != "INVIA":
            print("Invio annullato.")
            return 1
        sent, failed = send_messages(participants, settings)
        print(f"\nCompletato: {sent} inviate, {len(failed)} non riuscite.")
        return 1 if failed else 0
    except (OSError, ValueError, smtplib.SMTPException) as exc:
        parser.error(str(exc))


if __name__ == "__main__":
    raise SystemExit(main())
