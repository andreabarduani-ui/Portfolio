# -*- coding: utf-8 -*-
"""
Generatore di Scheda Finanziaria (A/B/C/D/E)
------------------------------------------------------
Crea un Excel con formule live: cambi un'ora o un costo e tutto si ricalcola,
inclusi i controlli dei vincoli del fondo (A<=35%, B>=40%, D<=25%).

Uso:  py financial_plan_builder.py
Output: Scheda_Finanziaria_Template.xlsx (stessa cartella)
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter

OUT = "Scheda_Finanziaria_Template.xlsx"

# ---------------------------------------------------------------------------
# Stili
# ---------------------------------------------------------------------------
TITLE_FONT   = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
HEAD_FONT    = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
SECT_FONT    = Font(name="Calibri", size=11, bold=True, color="1F3864")
BOLD         = Font(bold=True)
INPUT_FONT   = Font(color="7F6000", bold=True)

NAVY   = PatternFill("solid", fgColor="1F3864")   # intestazioni
BLUE   = PatternFill("solid", fgColor="2E5496")   # sezioni
YELLOW = PatternFill("solid", fgColor="FFF2CC")   # input utente
GREY   = PatternFill("solid", fgColor="D9D9D9")   # totali
GREEN  = PatternFill("solid", fgColor="C6EFCE")   # semaforo OK
RED    = PatternFill("solid", fgColor="FFC7CE")   # semaforo KO

thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

EUR = '#,##0.00\\ "€"'
ORE = '#,##0.00'
PCT = '0%'

CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT   = Alignment(horizontal="left",  vertical="center", wrap_text=True)
RIGHT  = Alignment(horizontal="right", vertical="center")

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Scheda"

# Larghezza colonne (A=Cod, B=Voce, C=Persona, D=€/h, E=Ore, F=Costo €)
for col, w in {"A":7, "B":46, "C":26, "D":10, "E":9, "F":14}.items():
    ws.column_dimensions[col].width = w

def setc(coord, value=None, font=None, fill=None, align=None, fmt=None, border=True):
    c = ws[coord]
    if value is not None: c.value = value
    if font:  c.font = font
    if fill:  c.fill = fill
    if align: c.alignment = align
    if fmt:   c.number_format = fmt
    if border: c.border = BORDER
    return c

# ===========================================================================
# TITOLO
# ===========================================================================
ws.merge_cells("A1:F1")
setc("A1", "SCHEDA FINANZIARIA — A/B/C/D/E", TITLE_FONT, NAVY, CENTER, border=False)
ws.row_dimensions[1].height = 26

# ===========================================================================
# ZONA 1 — PARAMETRI
# ===========================================================================
ws.merge_cells("A3:F3")
setc("A3", "1) PARAMETRI DI CALCOLO  (celle gialle = da compilare)", HEAD_FONT, BLUE, LEFT)

# Etichette colonna A, valori colonna B
def param(row, label, value, is_input=False, fmt=EUR, formula=False):
    setc(f"A{row}", label, BOLD if not is_input else None, align=LEFT)
    cell = setc(f"B{row}", value, INPUT_FONT if is_input else BOLD,
                YELLOW if is_input else None, RIGHT, fmt)
    return cell

param(4, "Target E — finanziamento desiderato", 10349, is_input=True)   # B4
param(5, "D % su C — costi indiretti (max 25%)", 0.25, is_input=True, fmt=PCT)  # B5
param(6, "C target (A+B)  = E / (1+D%)", "=B4/(1+B5)")                   # B6
param(7, "A massimo  (35% di E)", "=35%*B4")                            # B7
param(8, "B minimo  (40% di E)", "=40%*B4")                             # B8
param(9, "D calcolato target (C × D%)", "=B6*B5")                       # B9

E_TARGET = "$B$4"; D_PCT = "$B$5"; C_TARGET = "$B$6"
A_MAX = "$B$7";  B_MIN = "$B$8"

# ===========================================================================
# VOCI UFFICIALI DAL MANUALE (Tab. 11 e 12)
# ===========================================================================
VOCI_A = [
    ("A1","Progettazione esecutiva"),
    ("A2","Ricerche"),
    ("A3","Selezione - orientamento - bilancio delle competenze"),
    ("A4","Formazione formatori"),
    ("A5","Coordinamento e gestione"),
    ("A6","Monitoraggio e valutazione"),
    ("A7","Promozione e diffusione dei risultati"),
    ("A8","Fidejussione"),
    ("A9","Spese di viaggio"),
    ("A10","Altro (dettagliare analiticamente)"),
]
VOCI_B = [
    ("B1","Docenza"),
    ("B2","Tutoraggio"),
    ("B3","Sostegno all'utenza svantaggiata"),
    ("B4","Progettazione, elaborazione materiale didattico e FAD"),
    ("B5","Produzione, acquisto e distribuzione materiale didattico"),
    ("B6","Noleggi (mezzi e/o logistica)"),
    ("B7","Assicurazioni"),
    ("B8","Commissioni d'esame / certificazione competenze"),
    ("B9","Spese di viaggio"),
    ("B10","Altro (dettagliare analiticamente)"),
]

# Esempio prefisso (tratto dal BDG PROGETTO per validare: E deve fare 10.349 €)
PREFILL_A = {
    "A1": ("Pierpaolo Rossi", 21.27, 40),
    "A5": ("Coordinatore A", 25.54, 40),         # coordinamento
    "A5b":("Consulente B", 15.78, 31.4595),       # rendicontazione (2a riga A5)
    "A8": ("Polizza fidejussoria", 200, 1),
}
PREFILL_B = {
    "B2": ("Tutor C", 19.02, 82),       # tutor 2025
    "B2b":("Tutor C", 20.07, 26),       # tutor 2026
    "B1": ("Daniele Magli", 60, 30),                 # docenza
    "B10":("Coordinatore A - coord. d'aula", 25.54, 30),
    "B10b":("Monitor D - monitoraggio d'aula", 20.30, 52.3503),
}

def intestazione_tabella(title, row, limit_text):
    ws.merge_cells(f"A{row}:F{row}")
    setc(f"A{row}", f"{title}   ({limit_text})", HEAD_FONT, BLUE, LEFT)

def intestazione_colonne(row):
    for col, txt in zip("ABCDEF", ["Cod.","Voce di costo","Persona","€/h","Ore","Costo €"]):
        setc(f"{col}{row}", txt, BOLD, GREY, CENTER)

def riga_voce(row, cod, voce, persona=None, euh=None, ore=None):
    setc(f"A{row}", cod, align=CENTER)
    setc(f"B{row}", voce, align=LEFT)
    # C, D, E input (gialli se vuoti o sempre input)
    pc = setc(f"C{row}", persona, INPUT_FONT if persona else None,
              YELLOW if persona is None else None, LEFT)
    # se persona fornita (esempio) non giallo; se vuoto -> giallo per compilare
    if persona is None:
        ws[f"C{row}"].fill = YELLOW
    setc(f"D{row}", euh, INPUT_FONT if euh else None,
         YELLOW if euh is None else None, RIGHT, fmt=EUR)
    setc(f"E{row}", ore, INPUT_FONT if ore else None,
         YELLOW if ore is None else None, RIGHT, fmt=ORE)
    # F = costo calcolato
    setc(f"F{row}", f'=IF(OR(D{row}="",E{row}=""),"",D{row}*E{row})',
         BOLD, None, RIGHT, fmt=EUR)

# --- Sezione A ---
A_TITLE, A_HDR, A_FIRST, A_LAST = 11, 12, 13, 22
A_TOTAL = A_LAST + 1   # 23
intestazione_tabella("2) A — COSTI DIRETTI propedeutici, accompagnamento e finali",
                     A_TITLE, "MAX 35% di E")
intestazione_colonne(A_HDR)
extra_a = {"A5b": "A5", "B10b": "B10"}  # codici per righe aggiuntive
for i, (cod, voce) in enumerate(VOCI_A):
    r = A_FIRST + i
    pf = PREFILL_A.get(cod)
    if pf:
        riga_voce(r, cod, voce, pf[0], pf[1], pf[2])
    else:
        riga_voce(r, cod, voce)
# riga aggiuntiva A5 (rendicontazione)
r = A_LAST  # usa l'ultima riga libera per la 2a voce A5 -> la sovrascrivo su A10? no.
# Invece aggiungo la rendicontazione sulla riga A10 se libera: ma A10 e' prevista.
# Soluzione semplice: sovrascrivo A10 con la rendicontazione (esempio PROGETTO non usa A10).
pf = PREFILL_A.get("A5b")
riga_voce(A_LAST, "A5", "Coordinamento e gestione (Rendicontazione)", pf[0], pf[1], pf[2])

# Totale A
setc(f"A{A_TOTAL}", "", fill=GREY)
setc(f"B{A_TOTAL}", "TOTALE A", BOLD, GREY, LEFT)
for col in "CDE": setc(f"{col}{A_TOTAL}", "", fill=GREY)
setc(f"F{A_TOTAL}", f"=SUM(F{A_FIRST}:F{A_LAST})", BOLD, GREY, RIGHT, fmt=EUR)

# --- Sezione B ---
B_TITLE, B_HDR, B_FIRST, B_LAST = 24, 25, 26, 35
B_TOTAL = B_LAST + 1   # 36
intestazione_tabella("3) B — REALIZZAZIONE ATTIVITA' FORMATIVE (docenza, tutoraggio, ...)",
                     B_TITLE, "MIN 40% di E")
intestazione_colonne(B_HDR)
# prima tutte le righe vuote, poi sovrascrivo con i dati PROGETTO
for i, (cod, voce) in enumerate(VOCI_B):
    r = B_FIRST + i
    riga_voce(r, cod, voce)
def put(row, cod, voce, persona, euh, ore):
    riga_voce(row, cod, voce, persona, euh, ore)
put(B_FIRST+0, "B1", "Docenza", "Daniele Magli", 60, 30)
put(B_FIRST+1, "B2", "Tutoraggio (2025)", "Tutor C", 19.02, 82)
put(B_FIRST+2, "B2", "Tutoraggio (2026)", "Tutor C", 20.07, 26)
put(B_FIRST+8, "B10", "Altro - Coordinamento d'aula", "Coordinatore A", 25.54, 30)
put(B_FIRST+9, "B10", "Altro - Monitoraggio d'aula", "Monitor D", 20.30, 52.3503)

# Totale B
setc(f"A{B_TOTAL}", "", fill=GREY)
setc(f"B{B_TOTAL}", "TOTALE B", BOLD, GREY, LEFT)
for col in "CDE": setc(f"{col}{B_TOTAL}", "", fill=GREY)
setc(f"F{B_TOTAL}", f"=SUM(F{B_FIRST}:F{B_LAST})", BOLD, GREY, RIGHT, fmt=EUR)

A_TOT = f"F{A_TOTAL}"; B_TOT = f"F{B_TOTAL}"

# ===========================================================================
# ZONA 4 — RIEPILOGO + SEMAFORI
# ===========================================================================
R0 = 38
ws.merge_cells(f"A{R0}:F{R0}")
setc(f"A{R0}", "4) RIEPILOGO E CONTROLLI VINCOLI", HEAD_FONT, BLUE, LEFT)

def riep(row, label, formula, fmt=EUR, highlight=False):
    setc(f"A{row}", label, BOLD, align=LEFT)
    cell = setc(f"B{row}", formula, BOLD, GREY if highlight else None, RIGHT, fmt)
    return cell

C_ROW   = R0+1   # 39
D_ROW   = R0+2   # 40
EC_ROW  = R0+3   # 41
ET_ROW  = R0+4   # 42
DEL_ROW = R0+5   # 43
GOAL    = R0+6   # 44
CA      = R0+7   # 45
CB      = R0+8   # 46
CD      = R0+9   # 47

riep(C_ROW,  "C = A + B  (totale costi diretti)", f"={A_TOT}+{B_TOT}")
riep(D_ROW,  "D = forfait (C × D%)", f"=B{C_ROW}*{D_PCT}")
riep(EC_ROW, "E calcolato  (C + D)", f"=B{C_ROW}+B{D_ROW}", highlight=True)
riep(ET_ROW, "E target  (riferimento)", f"={E_TARGET}")
riep(DEL_ROW,"Δ  (E calcolato − E target)", f"=B{EC_ROW}-B{ET_ROW}")
# Stato obiettivo
setc(f"A{GOAL}", "Stato obiettivo", BOLD, align=LEFT)
setc(f"B{GOAL}",
     f'=IF(ABS(B{DEL_ROW})<0.5,"✓ RAGGIUNTO","Δ = "&TEXT(B{DEL_ROW},"0.00")&" €")',
     BOLD, None, RIGHT)
# Controlli vincoli
def chk(row, label, formula):
    setc(f"A{row}", label, BOLD, align=LEFT)
    setc(f"B{row}", formula, BOLD, None, CENTER)
chk(CA, "Vincolo A ≤ 35% di E",  f'=IF({A_TOT}<={A_MAX},"OK","SUPERATO")')
chk(CB, "Vincolo B ≥ 40% di E",  f'=IF({B_TOT}>={B_MIN},"OK","SOTTO MINIMO")')
chk(CD, "Vincolo D ≤ 25% di C",  f'=IF(B{D_ROW}<=B{C_ROW}*0.25,"OK","SUPERATO")')

# Formattazione condizionale (semafori) sui controlli e sull'obiettivo
for row in (CA, CB, CD):
    ws.conditional_formatting.add(f"B{row}",
        CellIsRule(operator="equal", formula=['"OK"'], fill=GREEN))
    ws.conditional_formatting.add(f"B{row}",
        CellIsRule(operator="notEqual", formula=['"OK"'], fill=RED))
ws.conditional_formatting.add(f"B{GOAL}",
    CellIsRule(operator="containsText", formula=['"RAGGIUNTO"'], fill=GREEN))
# Goal rosso se non contiene RAGGIUNTO
ws.conditional_formatting.add(f"B{GOAL}",
    CellIsRule(operator="notContains", formula=['"RAGGIUNTO"'], fill=RED))

# ===========================================================================
# ZONA 5 — SIMULATORE ORE NECESSARIE
# ===========================================================================
S0 = 49
ws.merge_cells(f"A{S0}:F{S0}")
setc(f"A{S0}", "5) SIMULATORE — quante ore mancano al target?", HEAD_FONT, BLUE, LEFT)

setc(f"A{S0+1}", "Residuo al target  (C_target − C_attuale)",
     BOLD, align=LEFT)
setc(f"B{S0+1}", f"={C_TARGET}-B{C_ROW}", BOLD, GREY, RIGHT, fmt=EUR)

setc(f"A{S0+2}", "Costo orario di riferimento  (a chi vuoi aggiungere ore)",
     BOLD, align=LEFT)
setc(f"B{S0+2}", 60, INPUT_FONT, YELLOW, RIGHT, fmt=EUR)

setc(f"A{S0+3}", "ORE necessarie per chiudere il target  (residuo ÷ €/h)",
     BOLD, align=LEFT)
setc(f"B{S0+3}", f'=IF(B{S0+2}=0,"",B{S0+1}/B{S0+2})',
     Font(bold=True, color="1F3864"), YELLOW, RIGHT, fmt=ORE)

# Note d'uso
N0 = S0 + 5   # 54
ws.merge_cells(f"A{N0}:F{N0}")
setc(f"A{N0}", "NOTE D'USO", SECT_FONT, fill=GREY, align=LEFT)
note = [
    "• Celle GIALLE = input da compilare. Tutto il resto è calcolato in automatico.",
    "• Inserisci persona, €/h e ore nelle tabelle A e B: il costo di riga e i totali si aggiornano da soli.",
    "• L'obiettivo è far coincidere 'E calcolato' con 'E target' (Δ = 0 → RAGGIUNTO).",
    "• I 3 semafori controllano i vincoli del fondo: A≤35%, B≥40%, D≤25%. Se uno è rosso, riequilibra.",
    "• Il simulatore (zona 5) ti dice quante ore aggiungere a una persona per chiudere il residuo al target.",
    "• Le voci A8/Fidejussione sono costi fissi: usa ore=1.",
]
for i, t in enumerate(note):
    ws.merge_cells(f"A{N0+1+i}:F{N0+1+i}")
    setc(f"A{N0+1+i}", t, align=LEFT, border=False)

# Blocca riquadro sotto i parametri
ws.sheet_view.showGridLines = False
ws.freeze_panes = "A11"

wb.save(OUT)
print("Generato:", OUT)
