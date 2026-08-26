# -*- coding: utf-8 -*-
"""
Adds a 'Calcolatore' sheet to the BDG_AUTOMATICO.xlsx file,
linking to the totals of the existing 'BDG PROGETTO' sheet.
Does NOT modify the original BDG PROGETTO sheet: it is left unchanged.
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import CellIsRule

FILE = "BDG_AUTOMATICO.xlsx"

# Styles
TITLE = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
HEAD  = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
BOLD  = Font(bold=True)
INPUT = Font(color="7F6000", bold=True)
BIG   = Font(bold=True, size=12, color="1F3864")
NAVY  = PatternFill("solid", fgColor="1F3864")
BLUE  = PatternFill("solid", fgColor="2E5496")
YEL   = PatternFill("solid", fgColor="FFF2CC")
GREY  = PatternFill("solid", fgColor="D9D9D9")
GREEN = PatternFill("solid", fgColor="C6EFCE")
RED   = PatternFill("solid", fgColor="FFC7CE")
thin  = Side(style="thin", color="BFBFBF")
BORD  = Border(left=thin, right=thin, top=thin, bottom=thin)
EUR = '#,##0.00\\ "€"'
PCT = '0%'
CEN = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEF = Alignment(horizontal="left",  vertical="center", wrap_text=True)
RIG = Alignment(horizontal="right", vertical="center")

wb = openpyxl.load_workbook(FILE)
# Remove any old Calcolatore sheet
if "Calcolatore" in wb.sheetnames:
    del wb["Calcolatore"]
ws = wb.create_sheet("Calcolatore", 0)   # in first position

ws.sheet_view.showGridLines = False
for col, w in {"A":4, "B":42, "C":16, "D":16, "E":16, "F":4, "G":16}.items():
    ws.column_dimensions[col].width = w

def S(coord):
    """Reference to the BDG PROGETTO sheet."""
    return f"'BDG PROGETTO'!{coord}"

def put(coord, value=None, font=None, fill=None, align=None, fmt=None, border=True):
    c = ws[coord]
    if value is not None: c.value = value
    if font:  c.font = font
    if fill:  c.fill = fill
    if align: c.alignment = align
    if fmt:   c.number_format = fmt
    if border: c.border = BORD
    return c

# Title
ws.merge_cells("B2:G2")
put("B2", "CALCOLATORE SCHEDA FINANZIARIA", TITLE, NAVY, CEN, border=False)
ws.row_dimensions[2].height = 28

# --- Parameters ---
ws.merge_cells("B4:G4")
put("B4", "PARAMETRI  (celle gialle = da compilare)", HEAD, BLUE, LEF)
put("B5", "Target E — finanziamento da raggiungere", font=BOLD, align=LEF)
put("C5", 10349, font=INPUT, fill=YEL, align=RIG, fmt=EUR)   # manual input
TARGET = "$C$5"
put("B6", "D % su C — costi indiretti (max 25%)", font=BOLD, align=LEF)
put("C6", 0.25, font=INPUT, fill=YEL, align=RIG, fmt=PCT)
DPCT = "$C$6"
put("B7", "C target (A+B) = E / (1+D%)", font=BOLD, align=LEF)
put("C7", f"={TARGET}/(1+{DPCT})", font=BOLD, fill=GREY, align=RIG, fmt=EUR)
put("B8", "A massimo (35% di E)", font=BOLD, align=LEF)
put("C8", f"=35%*{TARGET}", font=BOLD, fill=GREY, align=RIG, fmt=EUR)
put("B9", "B minimo (40% di E)", font=BOLD, align=LEF)
put("C9", f"=40%*{TARGET}", font=BOLD, fill=GREY, align=RIG, fmt=EUR)

# --- Current status (from BDG PROGETTO) ---
ws.merge_cells("B11:G11")
put("B11", "STATO ATTUALE DELLA SCHEDA  (dal foglio BDG PROGETTO)", HEAD, BLUE, LEF)
put("B12", "Totale A  (dal BDG PROGETTO, cella M16)", font=BOLD, align=LEF)
put("C12", f"={S('M16')}", font=BOLD, align=RIG, fmt=EUR)
put("B13", "Totale B  (dal BDG PROGETTO, cella M36)", font=BOLD, align=LEF)
put("C13", f"={S('M36')}", font=BOLD, align=RIG, fmt=EUR)
put("B14", "C = A + B", font=BOLD, align=LEF)
put("C14", "=C12+C13", font=BOLD, fill=GREY, align=RIG, fmt=EUR)
put("B15", "D = forfait (C x D%)", font=BOLD, align=LEF)
put("C15", f"=C14*{DPCT}", font=BOLD, fill=GREY, align=RIG, fmt=EUR)
put("B16", "E calcolato (C + D)", font=BIG, align=LEF)
put("C16", "=C14+C15", font=BIG, fill=YEL, align=RIG, fmt=EUR)
ECALC = "$C$16"

# --- Goal ---
put("B18", "E target (riferimento)", font=BOLD, align=LEF)
put("C18", f"={TARGET}", font=BOLD, align=RIG, fmt=EUR)
put("B19", "Differenza (E calcolato - E target)", font=BOLD, align=LEF)
put("C19", f"={ECALC}-{TARGET}", font=BOLD, align=RIG, fmt=EUR)
put("B20", "STATO OBIETTIVO", font=BOLD, align=LEF)
put("C20", f'=IF(ABS(C19)<0.5,"✓ RAGGIUNTO","Δ "&TEXT(C19,"0.00")&" €")',
     font=BOLD, align=CEN)

# --- Constraint status lights ---
ws.merge_cells("B22:G22")
put("B22", "CONTROLLI VINCOLI DEL FONDO", HEAD, BLUE, LEF)
put("B23", "A ≤ 35% di E", font=BOLD, align=LEF)
put("C23", f'=IF(C12<=C8,"OK ✓","SUPERATO")', font=BOLD, align=CEN)
put("B24", "B ≥ 40% di E", font=BOLD, align=LEF)
put("C24", f'=IF(C13>=C9,"OK ✓","SOTTO MINIMO")', font=BOLD, align=CEN)
put("B25", "D ≤ 25% di C", font=BOLD, align=LEF)
put("C25", f'=IF(C15<=C14*0.25,"OK ✓","SUPERATO")', font=BOLD, align=CEN)
for r in (23,24,25):
    ws.conditional_formatting.add(f"C{r}",
        CellIsRule(operator="equal", formula=['"OK ✓"'], fill=GREEN))
    ws.conditional_formatting.add(f"C{r}",
        CellIsRule(operator="notEqual", formula=['"OK ✓"'], fill=RED))
ws.conditional_formatting.add("C20",
    CellIsRule(operator="containsText", formula=['"RAGGIUNTO"'], fill=GREEN))
ws.conditional_formatting.add("C20",
    CellIsRule(operator="notContains", formula=['"RAGGIUNTO"'], fill=RED))

# --- Hours simulator ---
ws.merge_cells("B27:G27")
put("B27", "SIMULATORE — quante ore mancano al target?", HEAD, BLUE, LEF)
put("B28", "Residuo al target (C target - C attuale)", font=BOLD, align=LEF)
put("C28", f"=C7-C14", font=BOLD, fill=GREY, align=RIG, fmt=EUR)
put("B29", "Costo orario di riferimento (chi deve fare ore)", font=BOLD, align=LEF)
put("C29", 60, font=INPUT, fill=YEL, align=RIG, fmt=EUR)
put("B30", "ORE necessarie per chiudere = residuo / costo orario",
     font=BOLD, align=LEF)
put("C30", '=IF(C29=0,"",C28/C29)', font=BIG, fill=YEL, align=RIG, fmt='#,##0.00 "h"')

# --- Notes ---
ws.merge_cells("B32:G32")
put("B32", "COME SI USA", Font(bold=True, color="1F3864"), GREY, LEF)
note = [
 "1. In C5 scrivi il finanziamento target (es. 10.000). Tutto il resto si ricalcola.",
 "2. Le ore le inserisci nel foglio 'BDG PROGETTO' (colonna L): qui vedi i totali e i controlli.",
 "3. Obiettivo: 'E calcolato' (C16) deve coincidere con il target → STATO = RAGGIUNTO.",
 "4. I 3 semafori devono restare verdi: A≤35%, B≥40%, D≤25%.",
 "5. Simulatore: in C29 metti il costo orario di chi puo' fare altre ore; C30 ti dice quante servono.",
]
for i, t in enumerate(note):
    ws.merge_cells(f"B{33+i}:G{33+i}")
    put(f"B{33+i}", t, align=LEF, border=False)

ws.freeze_panes = "A4"
wb.save(FILE)
print("Calcolatore sheet added to:", FILE)
print("Sheets now present:", wb.sheetnames)
