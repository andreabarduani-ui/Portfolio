# -*- coding: utf-8 -*-
"""
timesheet_generator.py
=========================
Automatically fills in the timesheets (Allegato A.7) of the PROGRAMMA_A/PROGRAMMA_B
project from:
  - Lista Tirocini.xlsx     -> trainee records
  - Timesheet <tutor>.xlsx  -> tutoring hours (2 files: Tutor A + Tutor B)
  - Timesheet_template.xlsx  -> monthly template

For each trainee with activity in the reporting year it generates:
  1) one Excel file per month  (Timesheet_<Nome>_<AAAA-MM>.xlsx)
  2) a single Excel file with one sheet per month
                                 (Timesheet_<Nome>_UNICO_<AAAA>.xlsx)

The template is extended inline with a new "Denominazione Progetto" row
in position 5; the file itself is not modified (a "v2" copy is created).

The Programma_A/Programma_B distinction is managed via config_progetti.yaml.

Usage:
    py timesheet_generator.py            # runs everything
    py timesheet_generator.py --dry-run  # extraction + report only, no files
    py timesheet_generator.py --solo-unici   # generates only the single files

Dependencies:  openpyxl, pyyaml
"""
import argparse
import calendar
import datetime
import os
import re
import shutil
import sys
from collections import defaultdict
from copy import copy

import openpyxl
import yaml
from openpyxl.utils import get_column_letter

# ============================================================
# Constants
# ============================================================
NOMI_MESI = ['', 'Gennaio', 'Febbraio', 'Marzo', 'Aprile', 'Maggio', 'Giugno',
             'Luglio', 'Agosto', 'Settembre', 'Ottobre', 'Novembre', 'Dicembre']
GIORNI_SETT = ['Lun', 'Mar', 'Mer', 'Gio', 'Ven', 'Sab', 'Dom']
CONFIG_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                           'config_progetti.yaml')


# ============================================================
# Utility
# ============================================================
def norm(s):
    """Normalizes a string: lowercase, multiple spaces -> single, nbsp removed."""
    if s is None:
        return ''
    s = str(s).replace('\xa0', ' ')
    return re.sub(r'\s+', ' ', s).strip().lower()


def safe_name(s):
    """Removes characters that are invalid in Windows file/folder names."""
    for ch in ['<', '>', ':', '"', '/', '\\', '|', '?', '*']:
        s = s.replace(ch, '_')
    return s.strip()


def carica_config():
    """Loads config_progetti.yaml."""
    with open(CONFIG_PATH, 'r', encoding='utf-8') as f:
        cfg = yaml.safe_load(f)
    return cfg


# ============================================================
# Reading trainee records (Lista Tirocini)
# ============================================================
def leggi_anagrafiche(path_lista):
    """Returns a sorted list of dicts with the trainees' data."""
    wb = openpyxl.load_workbook(path_lista, data_only=True)
    ws = wb['Riepilogo tirocini attivati ']
    anag = []
    for r in range(15, 45):
        num = ws.cell(r, 2).value
        nome = ws.cell(r, 3).value
        if not nome:
            continue
        nome = re.sub(r'\s+', ' ', str(nome).replace('\xa0', ' ')).strip()
        anag.append({
            'num': num,
            'tirocinante': nome,
            'key': norm(nome),
            'cf': norm(ws.cell(r, 4).value),
            'so': re.sub(r'\s+', ' ', str(ws.cell(r, 5).value).replace('\xa0', ' ')).strip(),
            'inizio': ws.cell(r, 6).value,
            'fine': ws.cell(r, 7).value,
            'tutor': norm(ws.cell(r, 8).value),
        })
    return anag


# ============================================================
# Extracting hours from the tutor timesheets
# ============================================================
def normalizza_nome_tutor(nome_txt):
    """From 'Tirocinio Luca Novelli' / 'Tutoraggio Tirocinante X' -> 'luca novelli' / 'x'."""
    s = nome_txt
    for pref in ['Tutoraggio Tirocinante', 'Tirocinio', 'Tutor Tirocinio',
                 'Tutoraggio', 'Tutor']:
        s = re.sub(r'^\s*' + re.escape(pref), '', s, flags=re.IGNORECASE)
    return norm(s)


def estrai_da_tutor(file_path, anno):
    """Extracts hours per trainee from a tutor timesheet file.
    Returns dict: nome_norm -> {(mese, anno): {giorno: ore}}.
    Considers only rows with B='Programma_A' (or 'Programma_B') and C containing a name
    (excludes 'Segreteria')."""
    wb = openpyxl.load_workbook(file_path, data_only=True)
    ws = wb['imputazione ore']
    risultato = defaultdict(lambda: defaultdict(dict))
    mese_cor = anno_cor = None
    for r in range(3, 399):
        a = ws.cell(r, 1).value
        b = ws.cell(r, 2).value
        c = ws.cell(r, 3).value
        # update current month/year (the date is repeated on every row of the block)
        if a and isinstance(a, str):
            m = re.match(r'^\s*(\d{1,2})/(\d{4})\s*$', a)
            if m:
                mese_cor, anno_cor = int(m.group(1)), int(m.group(2))
        # Programma_A or Programma_B row with a trainee name
        if b and isinstance(b, str) and b.strip().lower() in ('programma_a', 'programma_b'):
            if c and isinstance(c, str):
                c_clean = c.strip()
                if 'segreteria' in c_clean.lower():
                    continue
                nome_norm = normalizza_nome_tutor(c_clean)
                if not nome_norm or len(nome_norm) < 3:
                    continue
                if mese_cor and anno_cor and anno_cor == anno:
                    for day in range(1, 32):
                        v = ws.cell(r, 810 + day).value
                        if isinstance(v, (int, float)) and v != 0:
                            risultato[nome_norm][(mese_cor, anno_cor)][day] = \
                                risultato[nome_norm][(mese_cor, anno_cor)].get(day, 0) + v
    return risultato


def estrai_ore_tutor(paths_tutor, anno):
    """Merges the extractions from multiple tutor files."""
    ore_all = defaultdict(lambda: defaultdict(dict))
    for p in paths_tutor:
        parziale = estrai_da_tutor(p, anno)
        for nome, mesi in parziale.items():
            for k, giorni in mesi.items():
                for d, val in giorni.items():
                    ore_all[nome][k][d] = ore_all[nome][k].get(d, 0) + val
    return ore_all


# ============================================================
# Matching Lista names <-> tutor labels
# ============================================================
def match_fuzzy(nl, nt):
    """Match between the list name (nl) and the tutor name (nt), both normalized."""
    if nl == nt:
        return 'ESATTO'
    parole_l = nl.split()
    parole_t = nt.split()
    if len(parole_l) >= 2 and ' '.join([parole_l[0], parole_l[-1]]) == nt:
        return 'NOME_COGNOME'
    if parole_l and len(parole_l[-1]) >= 4 and parole_l[-1] in parole_t:
        return 'COGNOME_OK'
    if nl in nt or nt in nl:
        return 'INCLUSIONE'
    return None


def costruisci_corrispondenze(anag, ore_all):
    """For each trainee, finds the corresponding tutor name."""
    rank_order = {'ESATTO': 4, 'NOME_COGNOME': 3, 'INCLUSIONE': 2, 'COGNOME_OK': 1}
    # known manual mappings (name variants)
    mappa_manuale = {
        'teresita de jesus saenz villarreal': 'teresita saenz villareal',
    }
    corr = {}
    report = []
    for info in anag:
        nk = info['key']
        if nk in mappa_manuale:
            corr[nk] = mappa_manuale[nk]
            report.append((info['tirocinante'], mappa_manuale[nk], 'MANUALE'))
            continue
        best, best_rank, best_type = None, -1, None
        for nt in ore_all.keys():
            m = match_fuzzy(nk, nt)
            if m:
                rk = rank_order.get(m, 0)
                if rk > best_rank:
                    best, best_rank, best_type = nt, rk, m
        if best:
            corr[nk] = best
            report.append((info['tirocinante'], best, best_type))
        else:
            report.append((info['tirocinante'], None, 'NESSUN_MATCH'))
    return corr, report


# ============================================================
# Template v2 (Denominazione Progetto row)
# ============================================================
def crea_template_v2(path_template_originale):
    """Creates a copy of the template with a new 'Denominazione Progetto'
    row in position 5. Manually handles merges and formulas (insert_rows does
    not update them automatically in openpyxl). Returns the v2 file path."""
    path_v2 = os.path.join(os.path.dirname(path_template_originale),
                           '_template_v2_tmp.xlsx')
    shutil.copy(path_template_originale, path_v2)
    wb = openpyxl.load_workbook(path_v2)
    ws = wb['Foglio1']

    # 1) remove existing merges
    for mr in list(ws.merged_cells.ranges):
        ws.unmerge_cells(str(mr))

    # 2) insert row at position 5
    ws.insert_rows(5, amount=1)

    # 3) re-insert merges at the correct positions
    for m in ['B4:AH4', 'C5:J5', 'C6:J6', 'C7:J7', 'C8:J8', 'C9:J9',
              'C10:J10', 'C11:J11', 'C12:J12', 'AH13:AH14']:
        ws.merge_cells(m)

    # 4) label + style of row 5 (style copied from row 6)
    ws.cell(5, 2).value = 'Denominazione Progetto'
    ws.cell(5, 3).value = ''
    for col in (2, 3):
        src_cell = ws.cell(6, col)
        dst_cell = ws.cell(5, col)
        dst_cell.font = copy(src_cell.font)
        dst_cell.fill = copy(src_cell.fill)
        dst_cell.alignment = copy(src_cell.alignment)
        dst_cell.border = copy(src_cell.border)
        dst_cell.number_format = src_cell.number_format
    ws.cell(5, 3).number_format = '@'
    ws.row_dimensions[5].height = ws.row_dimensions[6].height

    # 5) fix total formulas (activity rows now 15..23, total row 24)
    for r in range(15, 24):
        ws.cell(r, 34).value = '=SUM(C%d:AG%d)' % (r, r)
    for c in range(3, 34):
        cl = get_column_letter(c)
        ws.cell(24, c).value = '=SUM(%s15:%s23)' % (cl, cl)
    ws.cell(24, 34).value = '=SUM(C24:AG24)'

    wb.save(path_v2)
    return path_v2


def copia_foglio(ws_src, ws_dst):
    """Copies values, styles, merges, row/column dimensions from ws_src to ws_dst."""
    for row in ws_src.iter_rows():
        for cell in row:
            new = ws_dst.cell(row=cell.row, column=cell.column, value=cell.value)
            if cell.has_style:
                new.font = copy(cell.font)
                new.fill = copy(cell.fill)
                new.border = copy(cell.border)
                new.alignment = copy(cell.alignment)
                new.number_format = cell.number_format
                new.protection = copy(cell.protection)
    for mr in ws_src.merged_cells.ranges:
        ws_dst.merge_cells(str(mr))
    for k, dim in ws_src.column_dimensions.items():
        ws_dst.column_dimensions[k].width = dim.width
        ws_dst.column_dimensions[k].hidden = dim.hidden
    for k, dim in ws_src.row_dimensions.items():
        ws_dst.row_dimensions[k].height = dim.height


def compila_foglio_mese(ws, info, ore_mese, mese, anno, denominazione):
    """Fills a sheet (already populated from the v2 template) with records + hours."""
    giorni_mese = calendar.monthrange(anno, mese)[1]
    # records (template v2):
    # row 5=Denominazione Progetto, 6=CUP, 7=Cod.Progetto, 8=SO,
    # 9=Tutor, 10=Tirocinante, 11=start, 12=end
    ws.cell(5, 3).value = denominazione
    ws.cell(6, 3).value = 'F81J25000180009'
    ws.cell(7, 3).value = '24020DP000000025'
    ws.cell(8, 3).value = info['so']
    ws.cell(9, 3).value = info['tutor'].title()
    ws.cell(10, 3).value = info['tirocinante']
    ws.cell(11, 3).value = info['inizio']
    ws.cell(12, 3).value = info['fine']
    ws.cell(11, 3).number_format = 'DD/MM/YYYY'
    ws.cell(12, 3).number_format = 'DD/MM/YYYY'
    # weekdays (row 13) + day numbers (row 14)
    for col_idx in range(3, 34):  # C..AG -> giorni 1..31
        giorno = col_idx - 2
        if giorno <= giorni_mese:
            d = datetime.date(anno, mese, giorno)
            ws.cell(13, col_idx).value = GIORNI_SETT[d.weekday()]
            ws.cell(14, col_idx).value = giorno
        else:
            ws.cell(13, col_idx).value = None
            ws.cell(14, col_idx).value = None
    # hours row 15
    for col_idx in range(3, 34):
        giorno = col_idx - 2
        if giorno > giorni_mese:
            continue
        val = ore_mese.get(giorno, 0)
        ws.cell(15, col_idx).value = val if (val and val != 0) else None


def mesi_periodo(ini, fine, anno):
    """Returns the list of (month, year) between ini and fine and within the given year."""
    if not ini or not fine or not hasattr(ini, 'year'):
        return []
    if ini.year > anno:
        return []
    if fine.year < anno:
        return []
    start_m = ini.month if ini.year == anno else 1
    end_m = fine.month if fine.year == anno else 12
    return [(m, anno) for m in range(start_m, end_m + 1)]


# ============================================================
# File generation
# ============================================================
def denominazione_per_tirocinante(info, cfg):
    """Returns the project denomination (Programma_A/Programma_B) for the trainee."""
    programma_b_list = [norm(n) for n in (cfg.get('tirocinanti_programma_b') or [])]
    denom_map = cfg.get('denominazioni', {})
    if info['key'] in programma_b_list:
        return denom_map.get('programma_b', 'Programma B')
    return denom_map.get('programma_a', 'Programma A')


def genera_file_mensili_separati(info, mesi, ore_tir, denom, path_template_v2, out_dir):
    """Generates one .xlsx file per month."""
    n = 0
    cartella = os.path.join(out_dir, safe_name(info['tirocinante']))
    os.makedirs(cartella, exist_ok=True)
    for (mese, anno) in mesi:
        ore_mese = ore_tir.get((mese, anno), {})
        giorni_mese = calendar.monthrange(anno, mese)[1]
        nome_file = 'Timesheet_%s_%d-%02d.xlsx' % (
            safe_name(info['tirocinante'].replace(' ', '')), anno, mese)
        path_file = os.path.join(cartella, nome_file)
        shutil.copy(path_template_v2, path_file)
        wb = openpyxl.load_workbook(path_file)
        compila_foglio_mese(wb['Foglio1'], info, ore_mese, mese, anno, denom)
        wb.save(path_file)
        n += 1
    return n


def genera_file_unico(info, mesi, ore_tir, denom, path_template_v2, out_dir):
    """Generates a single .xlsx file with one sheet per month."""
    cartella = os.path.join(out_dir, safe_name(info['tirocinante']))
    os.makedirs(cartella, exist_ok=True)
    wb_tpl = openpyxl.load_workbook(path_template_v2)
    ws_tpl = wb_tpl['Foglio1']
    wb_out = openpyxl.Workbook()
    wb_out.remove(wb_out.active)
    for (mese, anno) in mesi:
        ore_mese = ore_tir.get((mese, anno), {})
        nome_foglio = '%s %d' % (NOMI_MESI[mese], anno)
        ws_new = wb_out.create_sheet(title=nome_foglio)
        copia_foglio(ws_tpl, ws_new)
        compila_foglio_mese(ws_new, info, ore_mese, mese, anno, denom)
    nome_file = 'Timesheet_%s_UNICO_%d.xlsx' % (
        safe_name(info['tirocinante'].replace(' ', '')), mesi[0][1])
    path_file = os.path.join(cartella, nome_file)
    wb_out.save(path_file)
    return path_file


# ============================================================
# Main
# ============================================================
def main():
    parser = argparse.ArgumentParser(
        description='Fills in the Programma_A/Programma_B timesheets from the tutor timesheets.')
    parser.add_argument('--dry-run', action='store_true',
                        help='Extraction and report only, generates no files.')
    parser.add_argument('--solo-unici', action='store_true',
                        help='Generates only the single files (not the separate monthly ones).')
    args = parser.parse_args()

    base_dir = os.path.dirname(os.path.abspath(__file__))
    cfg = carica_config()
    file_cfg = cfg['file']
    anno = cfg['anno_rendiconto']
    print('>> Configuration loaded. Reporting year: %d' % anno)

    # 1) trainee records
    path_lista = os.path.join(base_dir, file_cfg['lista_tirocini'])
    print('>> Reading trainee records: %s' % os.path.basename(path_lista))
    anag = leggi_anagrafiche(path_lista)
    print('   %d trainees found.' % len(anag))

    # 2) hour extraction
    path_tutor = [
        os.path.join(base_dir, file_cfg['tutor_a']),
        os.path.join(base_dir, file_cfg['tutor_b']),
    ]
    print('>> Extracting hours from the tutor timesheets...')
    ore_all = estrai_ore_tutor(path_tutor, anno)
    print('   %d trainee labels extracted from the tutors.' % len(ore_all))

    # 3) matches
    corr, report_corr = costruisci_corrispondenze(anag, ore_all)

    # matching report
    print('\n=== NAME MATCHES ===')
    n_match = n_no = 0
    for nome_disp, nt, tipo in report_corr:
        flag = 'OK' if nt else 'NO'
        if nt:
            n_match += 1
        else:
            n_no += 1
        print('  [%s] %s -> %s (%s)' % (flag, nome_disp, nt or '-', tipo))
    print('Matched: %d, unmatched: %d\n' % (n_match, n_no))

    if args.dry_run:
        print('>> DRY-RUN: no files generated.')
        # still print the hour summary for those who would have files
        for info in anag:
            mesi = mesi_periodo(info['inizio'], info['fine'], anno)
            if not mesi:
                continue
            ore_tir = ore_all.get(corr.get(info['key'], ''), {})
            denom = denominazione_per_tirocinante(info, cfg)
            tot = []
            for (m, a) in mesi:
                ore_m = ore_tir.get((m, a), {})
                tot.append('%02d: %s h' % (m, sum(ore_m.values())))
            print('  %s | %s | %d mesi | %s' %
                  (info['tirocinante'], denom, len(mesi), ', '.join(tot)))
        return

    # 4) create v2 template (with Denominazione Progetto row)
    print('>> Creating v2 template...')
    path_tpl_v2 = crea_template_v2(os.path.join(base_dir, file_cfg['template']))
    print('   Template v2: %s' % os.path.basename(path_tpl_v2))

    # 5) generate files
    out_dir = os.path.join(base_dir, file_cfg['cartella_output'])
    os.makedirs(out_dir, exist_ok=True)
    n_unici = n_separati = n_saltati = 0
    print('\n=== FILE GENERATION ===')
    for info in anag:
        mesi = mesi_periodo(info['inizio'], info['fine'], anno)
        if not mesi:
            print('  SKIP %s (N.%s) - no months in %d' %
                  (info['tirocinante'], info['num'], anno))
            n_saltati += 1
            continue
        ore_tir = ore_all.get(corr.get(info['key'], ''), {})
        denom = denominazione_per_tirocinante(info, cfg)

        if not args.solo_unici:
            n_separati += genera_file_mensili_separati(
                info, mesi, ore_tir, denom, path_tpl_v2, out_dir)
        genera_file_unico(info, mesi, ore_tir, denom, path_tpl_v2, out_dir)
        n_unici += 1
        tot_ore = sum(sum(ore_tir.get((m, a), {}).values()) for (m, a) in mesi)
        print('  OK %s - %d months - denom=%s - total hours: %s' %
              (info['tirocinante'], len(mesi), denom.split()[-1], tot_ore))

    # 6) cleanup temporary v2 template
    try:
        os.remove(path_tpl_v2)
    except OSError:
        pass

    print('\n=== SUMMARY ===')
    print('Trainees processed: %d' % n_unici)
    print('Trainees skipped (other year): %d' % n_saltati)
    if not args.solo_unici:
        print('Separate monthly files generated: %d' % n_separati)
    print('Single files generated: %d' % n_unici)
    print('Output: %s' % out_dir)


if __name__ == '__main__':
    main()
