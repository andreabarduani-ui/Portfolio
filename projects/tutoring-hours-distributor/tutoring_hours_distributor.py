# -*- coding: utf-8 -*-
"""
Tutoring hours distribution (internship programme) over multiple months — VERSION 3
=========================================================================

FIXES COMPARED TO VERSION 2 (bug "hours all concentrated in one month")
----------------------------------------------------------------
1) ALWAYS START FROM THE INTERNSHIP START DATE.
   In v2 the starting month was taken from where the hours were MARKED in
   the timesheet (min(mesi_grezzi)). If that month fell towards the END of
   the internship, the start->end window was extremely narrow and all the
   hours got compressed there (via RIDISTRIBUISCI_ECCEDENZE) instead of
   being spread over the months. Now it ALWAYS starts from the internship
   data_inizio.

2) AVAILABILITY-AWARE DISTRIBUTION.
   In v2 distribuisci_mesi assigned each month a fixed quota
   min(MAX_ORE_MESE, rimanenti) WITHOUT knowing whether that month had free
   slots: the occupancy check (occupazione_mese) happened only AFTERWARDS,
   in proponi_slot_mese. So a month already saturated with other activities
   still received the quota, the hours became "unplaceable" and did NOT
   migrate to the freer months. Now distribuisci_mesi truly places only
   what FITS in the free slots (by calling proponi_slot_mese during the
   distribution) and the flow continues on the following months.

MONTHLY CAP: MAX_ORE_MESE = 4 (unchanged).

WORKFLOW (unchanged, two passes):
   - 1st run: use the "Proposta" to MANUALLY correct the real timesheet
     (the script never modifies the original files).
   - Feel free to run it again after the correction to verify.

LIBRARY INSTALLATION (once only, from PowerShell):
    py -m pip install openpyxl python-docx --user
"""

import re
import calendar
import difflib
import unicodedata
from datetime import datetime, date
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter, column_index_from_string

# ============================================================
# CONFIG — adapt these values to the real structure of your files
# ============================================================

# --- TIMESHEET FILE (opened read-only) ---
TIMESHEET_PATH = "timesheet_tutor.xlsx"
TIMESHEET_SHEET = "imputazione ore"

TS_HEADER_GIORNI_ROW = 1   # row with the day numbers (1, 2, 3, ...)
TS_HEADER_SLOT_ROW = 2     # row with the time slots ("08-09", "09-10", ...)
TS_DATA_START_ROW = 3      # first data row (below the headers)

COL_TS_MESE = "A"       # column "MESE E ANNO"
COL_TS_PROGETTO = "B"   # column "Progetto"
COL_TS_MANSIONE = "C"   # column "Mansione"
COL_TS_TOTALE = "D"     # column "TOTALE H MENSILI"

# Names extracted from "Tirocinio ..." that are NOT trainees (generic rows
# like "Tutoraggio Tirocinio Ente Attuatore"): silently ignored.
NOMI_DA_IGNORARE = {"ente attuatore"}

# Similarity threshold (0-1) for fuzzy name matching.
SOGLIA_FUZZY = 0.65

# Time slots NOT to be used for proposals: excluded ranges [start, end).
ORARI_ESCLUSI = [(8, 9), (13, 14), (18, 20)]

ESCLUDI_WEEKEND = True
UNITA_BASE = 0.5  # each half cell is worth 0.5 hours

# --- INTERNSHIPS FILE (read-only; the "grezzo" is a modified copy) ---
FILE_TIROCINI_PATH = "file_tirocini.xlsx"
FILE_TIROCINI_SHEET = None

TIROCINI_HEADER_ROW = 1
TIROCINI_DATA_START_ROW = 2

COL_NOME_TIROCINANTE = "A"
COL_DATA_INIZIO = "C"
COL_DATA_FINE = "D"
COL_STATO = "E"

PRIMA_COLONNA_MESE = "F"
ULTIMA_COLONNA_MESE = "T"
ANNO_PRIMA_COLONNA = 2025
MESE_PRIMA_COLONNA = 7

MAX_ORE_MESE = 4
RISPETTA_DATA_FINE = True

# If True and the available months are not enough with MAX_ORE_MESE, the
# leftover hours are still placed in the available months (exceeding the
# cap), with a warning in the report. If False, they stay "undistributed".
RIDISTRIBUISCI_ECCEDENZE = True

# --- OUTPUT (only two files) ---
CARTELLA_OUTPUT = "output"
OUTPUT_PROPOSTA_XLSX = CARTELLA_OUTPUT + r"\Proposta ore tutoraggio.xlsx"
OUTPUT_REPORT_DOCX = CARTELLA_OUTPUT + r"\Report ore tutoraggio.docx"

# "Grezzo" file: copy of the internships file with the month-by-month hours
# EXACTLY AS THEY ARE NOW in the timesheet (on first run: all concentrated
# in a single month). Set to False if you don't need it.
GENERA_FILE_GREZZO = True
OUTPUT_PATH_GREZZO = CARTELLA_OUTPUT + r"\File tirocini - grezzo.xlsx"

# ============================================================
# CONSTANTS / REGEX
# ============================================================

RE_TIROCINIO = re.compile(r"tirocinio[\s:\-]+(.+)", re.IGNORECASE)
RE_SLOT = re.compile(r"^(\d{1,2})\s*-\s*(\d{1,2})$")
RE_DATA_NEL_TESTO = re.compile(r"(\d{1,2})[/\-.](\d{1,2})[/\-.](\d{2,4})")

# ============================================================
# GENERIC HELPER FUNCTIONS
# ============================================================

def normalizza_testo(s):
    if s is None:
        return ""
    s = str(s).strip().lower()
    s = unicodedata.normalize("NFKD", s)
    s = "".join(c for c in s if not unicodedata.combining(c))
    s = re.sub(r"\s+", " ", s)
    return s


def parse_mese_anno(valore):
    if isinstance(valore, datetime):
        return valore.year, valore.month
    s = str(valore).strip()
    m = re.match(r"^(\d{1,2})[/\-](\d{4})$", s)
    if m:
        return int(m.group(2)), int(m.group(1))
    raise ValueError(f"formato mese/anno non riconosciuto: '{valore}'")


def parse_data_generica(valore):
    """Accepts datetime/date or text containing a dd/mm/yyyy date
    (e.g. 'proroga 16/4/2026', 'INTERROTTO 9/01/2026'). Returns date or None."""
    if isinstance(valore, datetime):
        return valore.date()
    if isinstance(valore, date):
        return valore
    if valore is None:
        return None
    m = RE_DATA_NEL_TESTO.search(str(valore))
    if not m:
        return None
    g, mese, anno = int(m.group(1)), int(m.group(2)), int(m.group(3))
    if anno < 100:
        anno += 2000
    try:
        return date(anno, mese, g)
    except ValueError:
        return None


def punteggio_similarita(a_norm, b_norm):
    ratio = difflib.SequenceMatcher(None, a_norm, b_norm).ratio()
    tokens_a, tokens_b = set(a_norm.split()), set(b_norm.split())
    overlap = (len(tokens_a & tokens_b) / min(len(tokens_a), len(tokens_b))
               if tokens_a and tokens_b else 0.0)
    return max(ratio, overlap)


def slot_escluso(slot_label):
    m = RE_SLOT.match(slot_label)
    if not m:
        return True
    h = int(m.group(1))
    return any(inizio <= h < fine for inizio, fine in ORARI_ESCLUSI)


# ============================================================
# AUTOMATIC DETECTION OF THE TIMESHEET TIME GRID
# ============================================================

def rileva_griglia(ws_ts):
    """Reads the time-slot row (TS_HEADER_SLOT_ROW) and reconstructs the
    real geometry of the grid:
      - first column of day 1;
      - list of the day's slots with, for each, HOW MANY columns it spans
        (in the real file: 2 columns of 0.5h each);
      - step (number of columns) between one day and the next;
      - maximum number of days present in the sheet.
    Returns a dictionary with this information."""
    colonne_slot = []  # (col_idx, label)
    for c in range(1, ws_ts.max_column + 1):
        v = ws_ts.cell(row=TS_HEADER_SLOT_ROW, column=c).value
        if v is None:
            continue
        s = str(v).strip()
        if RE_SLOT.match(s):
            colonne_slot.append((c, s))
    if not colonne_slot:
        raise RuntimeError(
            "Unable to detect the time grid: no header like '08-09' found "
            f"at row {TS_HEADER_SLOT_ROW} of the timesheet."
        )

    prima_col = colonne_slot[0][0]
    prima_label = colonne_slot[0][1]

    # Day 2 starts at the second occurrence of the first slot.
    successive = [c for c, lab in colonne_slot[1:] if lab == prima_label]
    passo_giorno = (successive[0] - prima_col) if successive else (ws_ts.max_column - prima_col + 1)

    # Slots of day 1, with the column width of each.
    fasce_giorno1 = [(c, lab) for c, lab in colonne_slot if c < prima_col + passo_giorno]
    slot_info = []  # [{label, offset, larghezza}]
    for i, (c, lab) in enumerate(fasce_giorno1):
        col_succ = fasce_giorno1[i + 1][0] if i + 1 < len(fasce_giorno1) else prima_col + passo_giorno
        slot_info.append({"label": lab, "offset": c - prima_col, "larghezza": col_succ - c})

    num_giorni = max(1, (ws_ts.max_column - prima_col + 1) // passo_giorno)

    return {
        "prima_col": prima_col,
        "passo_giorno": passo_giorno,
        "slot_info": slot_info,
        "num_giorni": num_giorni,
    }


def celle_slot(griglia, giorno, slot):
    """Returns the indices of ALL the columns (half cells) of a time slot
    on a given day."""
    base = griglia["prima_col"] + (giorno - 1) * griglia["passo_giorno"] + slot["offset"]
    return list(range(base, base + slot["larghezza"]))


# ============================================================
# TIMESHEET READING
# ============================================================

def estrai_tutoraggi(ws_ts, report):
    col_mese = column_index_from_string(COL_TS_MESE)
    col_mansione = column_index_from_string(COL_TS_MANSIONE)
    col_progetto = column_index_from_string(COL_TS_PROGETTO)
    col_totale = column_index_from_string(COL_TS_TOTALE)

    tutoraggi = {}
    for riga in range(TS_DATA_START_ROW, ws_ts.max_row + 1):
        mansione = ws_ts.cell(row=riga, column=col_mansione).value
        if mansione is None:
            continue
        m = RE_TIROCINIO.search(str(mansione))
        if not m:
            continue
        nome = re.sub(r"^[\s:\-]+", "", m.group(1).strip())
        nome = re.sub(r"\s*\([^)]*\)\s*$", "", nome).strip()
        if not nome:
            continue
        nome_norm = normalizza_testo(nome)
        if nome_norm in NOMI_DA_IGNORARE:
            continue

        mese_anno_val = ws_ts.cell(row=riga, column=col_mese).value
        try:
            anno, mese = parse_mese_anno(mese_anno_val)
        except ValueError as e:
            report.avviso(f"Riga {riga} del timesheet ignorata: {e}")
            continue

        ore = float(ws_ts.cell(row=riga, column=col_totale).value or 0)
        progetto = ws_ts.cell(row=riga, column=col_progetto).value
        d = tutoraggi.setdefault(nome_norm, {
            "nome_originale": nome, "righe": [],
            "mansione": str(mansione).strip(),
            "progetto": str(progetto).strip() if progetto else "",
        })
        d["righe"].append({"anno": anno, "mese": mese, "ore": ore})
        d.setdefault("righe_ts", set()).add(riga)
    return tutoraggi


def occupazione_mese(ws_ts, griglia, anno, mese, righe_da_ignorare):
    """Set of the half cells (day, col_assoluta) already occupied in the
    month, considering ALL the rows/projects of that month, EXCEPT the rows
    of the tutoring entries we are redistributing (righe_da_ignorare): those
    hours will be removed from the timesheet, so their cells are to be
    considered free."""
    col_mese = column_index_from_string(COL_TS_MESE)
    occupate = set()
    for riga in range(TS_DATA_START_ROW, ws_ts.max_row + 1):
        if riga in righe_da_ignorare:
            continue
        mese_val = ws_ts.cell(row=riga, column=col_mese).value
        if mese_val is None:
            continue
        try:
            r_anno, r_mese = parse_mese_anno(mese_val)
        except ValueError:
            continue
        if (r_anno, r_mese) != (anno, mese):
            continue
        ultimo_giorno = calendar.monthrange(anno, mese)[1]
        for g in range(1, min(ultimo_giorno, griglia["num_giorni"]) + 1):
            for slot in griglia["slot_info"]:
                for col in celle_slot(griglia, g, slot):
                    if ws_ts.cell(row=riga, column=col).value:
                        occupate.add((g, col))
    return occupate


# ============================================================
# INTERNSHIPS FILE
# ============================================================

def costruisci_mappa_mesi(ws_tir):
    col_start = column_index_from_string(PRIMA_COLONNA_MESE)
    col_end = column_index_from_string(ULTIMA_COLONNA_MESE)
    mappa, anno, mese = [], ANNO_PRIMA_COLONNA, MESE_PRIMA_COLONNA
    for col_idx in range(col_start, col_end + 1):
        mappa.append({"col_idx": col_idx, "anno": anno, "mese": mese})
        mese += 1
        if mese > 12:
            mese, anno = 1, anno + 1
    return mappa


def trova_riga_tirocinante(ws_tir, nome_norm):
    col_nome = column_index_from_string(COL_NOME_TIROCINANTE)
    candidati = []
    for riga in range(TIROCINI_DATA_START_ROW, ws_tir.max_row + 1):
        v = ws_tir.cell(row=riga, column=col_nome).value
        if v is None:
            continue
        nv = normalizza_testo(v)
        if nv == nome_norm:
            return riga, str(v).strip(), False, 1.0
        candidati.append((riga, str(v).strip(), nv))
    best, best_p = None, 0.0
    for riga, orig, nv in candidati:
        p = punteggio_similarita(nome_norm, nv)
        if p > best_p:
            best_p, best = p, (riga, orig)
    if best and best_p >= SOGLIA_FUZZY:
        return best[0], best[1], True, best_p
    return None, None, False, 0.0


# ============================================================
# DISTRIBUTION  (VERSION 3 — availability-aware)
# ============================================================

def distribuisci_mesi(ore_totali, mesi_finestra, capacita_mese):
    """Distributes ore_totali over the months of 'mesi_finestra' in an
    AVAILABILITY-AWARE way: hours are placed ONLY where there are actually
    free slots (checks occupazione_mese through the callback).

    1st pass: up to MAX_ORE_MESE per month.
    2nd pass (if RIDISTRIBUISCI_ECCEDENZE and hours remain): goes over the
    months again with no cap. If a month is saturated with other activities,
    fewer (or zero) hours go there and the flow continues on the following
    months.

    Args:
        ore_totali: hours to distribute.
        mesi_finestra: list of dicts {anno, mese} already limited to the
            internship start/end interval, in the order to fill.
        capacita_mese(anno, mese, tetto) -> (proposte, ore_mancanti):
            callback that places up to `tetto` hours in the month's free slots.
            NOTE: the callback internally mutates the occupancy set (cache),
            so slots already used in a month are not taken again.

    Returns (allocazioni, ore_non_distribuite, eccedenza_ridistribuita) with
    allocazioni = [{anno, mese, ore, proposte}] in month order.
    """
    allocazioni = []
    alloc_idx = {}            # (anno, mese) -> position in allocazioni
    rimanenti = ore_totali

    def colloca(m, tetto):
        proposte, _mancanti = capacita_mese(m["anno"], m["mese"], tetto)
        placed = len(proposte) * UNITA_BASE
        if placed <= 0:
            return 0.0
        key = (m["anno"], m["mese"])
        if key not in alloc_idx:
            alloc_idx[key] = len(allocazioni)
            allocazioni.append({"anno": m["anno"], "mese": m["mese"],
                                "ore": 0.0, "proposte": []})
        voce = allocazioni[alloc_idx[key]]
        voce["ore"] += placed
        voce["proposte"].extend(proposte)
        return placed

    # 1st pass: MAX_ORE_MESE cap, month by month.
    for m in mesi_finestra:
        if rimanenti <= 1e-9:
            break
        rimanenti -= colloca(m, min(MAX_ORE_MESE, rimanenti))

    # 2nd pass (excess): goes over everything again with no cap.
    eccedenza = 0.0
    if rimanenti > 1e-9 and RIDISTRIBUISCI_ECCEDENZE:
        eccedenza = rimanenti
        for m in mesi_finestra:
            if rimanenti <= 1e-9:
                break
            rimanenti -= colloca(m, rimanenti)

    return allocazioni, rimanenti, eccedenza


def proponi_slot_mese(occupate, griglia, anno, mese, ore_necessarie, giorno_min, giorno_max):
    """Proposes free half cells in the month. Returns (proposte, ore_mancanti):
    proposte = [(giorno, slot_label, col_assoluta)].

    NOTE: mutates the 'occupate' set (adding the proposed cells). This is
    intentional: it guarantees that multiple calls in the same month (even
    for different trainees) do not overlap, because 'occupate' is shared
    via cache."""
    ultimo_giorno = calendar.monthrange(anno, mese)[1]
    giorno_max = min(giorno_max, ultimo_giorno, griglia["num_giorni"])
    celle_necessarie = round(ore_necessarie / UNITA_BASE)
    proposte = []
    for g in range(giorno_min, giorno_max + 1):
        if celle_necessarie <= 0:
            break
        if ESCLUDI_WEEKEND and date(anno, mese, g).weekday() >= 5:
            continue
        for slot in griglia["slot_info"]:
            if celle_necessarie <= 0:
                break
            if slot_escluso(slot["label"]):
                continue
            for col in celle_slot(griglia, g, slot):
                if celle_necessarie <= 0:
                    break
                if (g, col) in occupate:
                    continue
                occupate.add((g, col))
                proposte.append((g, slot["label"], col))
                celle_necessarie -= 1
    return proposte, celle_necessarie * UNITA_BASE


def descrivi_proposte(proposte):
    """Groups the half cells into readable text: 'gg 03: 09-10, 10-11 (½)'."""
    per_giorno = {}
    for g, lab, _col in proposte:
        per_giorno.setdefault(g, {}).setdefault(lab, 0)
        per_giorno[g][lab] += 1
    parti = []
    for g in sorted(per_giorno):
        fasce = []
        for lab in sorted(per_giorno[g]):
            n = per_giorno[g][lab]
            fasce.append(lab if n >= 2 else f"{lab} (½ ora)")
        parti.append(f"giorno {g:02d}: " + ", ".join(fasce))
    return parti


# ============================================================
# REPORT (collector + Word or txt writing)
# ============================================================

class Report:
    def __init__(self):
        self.sezioni = []     # per trainee: dict
        self.avvisi = []      # general strings
        self.errori = []

    def avviso(self, testo):
        self.avvisi.append(testo)

    def errore(self, testo):
        self.errori.append(testo)


def scrivi_report_docx(report, percorso, data_riferimento):
    try:
        from docx import Document
        from docx.shared import Pt, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH
    except ImportError:
        return False

    doc = Document()
    stile = doc.styles["Normal"]
    stile.font.name = "Calibri"
    stile.font.size = Pt(12)

    titolo = doc.add_heading("Distribuzione ore di tutoraggio", level=0)
    p = doc.add_paragraph(f"Report generato il {data_riferimento.strftime('%d/%m/%Y')}")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT

    doc.add_heading("Come usare questo report", level=1)
    doc.add_paragraph(
        "Per ogni tirocinante trovi la proposta di distribuzione delle ore, "
        "mese per mese, con i giorni e le fasce orarie da segnare nel timesheet. "
        "Le stesse ore sono già pronte, cella per cella, nel file Excel "
        "\u201cProposta ore tutoraggio.xlsx\u201d: basta ricopiare i valori 0,5 "
        "nelle stesse posizioni del timesheet."
    )

    n_ok = len(report.sezioni)
    doc.add_heading("Riepilogo", level=1)
    doc.add_paragraph(f"Tirocinanti elaborati: {n_ok}")
    doc.add_paragraph(f"Punti da controllare a mano: {len(report.avvisi)}")
    doc.add_paragraph(f"Errori: {len(report.errori)}")

    if report.errori or report.avvisi:
        doc.add_heading("Da controllare a mano", level=1)
        for e in report.errori:
            par = doc.add_paragraph(style="List Bullet")
            run = par.add_run("ERRORE: " + e)
            run.bold = True
            run.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
        for a in report.avvisi:
            par = doc.add_paragraph(style="List Bullet")
            run = par.add_run(a)
            run.font.color.rgb = RGBColor(0x99, 0x66, 0x00)

    doc.add_heading("Proposta per tirocinante", level=1)
    for sez in report.sezioni:
        doc.add_heading(sez["nome"], level=2)
        doc.add_paragraph(sez["intestazione"])
        for nota in sez["note"]:
            par = doc.add_paragraph(style="List Bullet")
            run = par.add_run(nota)
            run.font.color.rgb = RGBColor(0x99, 0x66, 0x00)

        tab = doc.add_table(rows=1, cols=3)
        tab.style = "Light Grid Accent 1"
        hdr = tab.rows[0].cells
        hdr[0].text = "Mese"
        hdr[1].text = "Ore"
        hdr[2].text = "Giorni e fasce orarie proposte"
        for r in sez["righe_mese"]:
            cells = tab.add_row().cells
            cells[0].text = r["mese"]
            cells[1].text = r["ore"]
            cells[2].text = "\n".join(r["dettaglio"])
        doc.add_paragraph()

    doc.save(percorso)
    return True


def scrivi_report_txt(report, percorso, data_riferimento):
    righe = ["REPORT DISTRIBUZIONE ORE TUTORAGGIO",
             f"Generato il {data_riferimento.strftime('%d/%m/%Y')}", "=" * 70, ""]
    if report.errori or report.avvisi:
        righe.append("DA CONTROLLARE A MANO:")
        for e in report.errori:
            righe.append(f"  [ERRORE] {e}")
        for a in report.avvisi:
            righe.append(f"  [ATTENZIONE] {a}")
        righe.append("")
    for sez in report.sezioni:
        righe.append(sez["nome"])
        righe.append("  " + sez["intestazione"])
        for nota in sez["note"]:
            righe.append(f"  [ATTENZIONE] {nota}")
        for r in sez["righe_mese"]:
            righe.append(f"    - {r['mese']} ({r['ore']}):")
            for d in r["dettaglio"]:
                righe.append(f"        {d}")
        righe.append("")
    righe.append("-" * 70)
    righe.append(f"Totale: {len(report.sezioni)} tirocinanti elaborati, "
                 f"{len(report.avvisi)} avvisi, {len(report.errori)} errori.")
    with open(percorso, "w", encoding="utf-8") as f:
        f.write("\n".join(righe) + "\n")


# ============================================================
# EXCEL "PROPOSTA" OUTPUT WITH LAYOUT IDENTICAL TO THE TIMESHEET
# ============================================================

def scrivi_proposta_xlsx(percorso, griglia, ws_ts, righe_proposta):
    """righe_proposta: list of dicts with
       {anno, mese, tirocinante, progetto, mansione, celle: [(giorno, col_assoluta)]}
    Creates a sheet with the same geometry as the timesheet: same columns,
    days in row 1, slots in row 2, 0.5 values highlighted in yellow."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Proposta (come timesheet)"

    grassetto = Font(bold=True)
    giallo = PatternFill(fill_type="solid", start_color="FFF2A9", end_color="FFF2A9")
    grigio = PatternFill(fill_type="solid", start_color="D9D9D9", end_color="D9D9D9")
    centro = Alignment(horizontal="center", vertical="center")
    bordo = Border(*[Side(style="thin")] * 4)

    # Fixed headers (copied from the timesheet)
    for col_lettera, testo in [(COL_TS_MESE, "MESE E ANNO"), (COL_TS_PROGETTO, "Progetto"),
                               (COL_TS_MANSIONE, "Mansione"), (COL_TS_TOTALE, "TOTALE H MENSILI")]:
        c = ws.cell(row=TS_HEADER_SLOT_ROW, column=column_index_from_string(col_lettera), value=testo)
        c.font, c.alignment = grassetto, centro

    # Grid headers: day numbers (row 1) and slots (row 2),
    # with the same columns as the original timesheet.
    max_giorni = griglia["num_giorni"]
    for g in range(1, max_giorni + 1):
        base = griglia["prima_col"] + (g - 1) * griglia["passo_giorno"]
        fine_giorno = base + griglia["passo_giorno"] - 1
        ws.merge_cells(start_row=TS_HEADER_GIORNI_ROW, start_column=base,
                       end_row=TS_HEADER_GIORNI_ROW, end_column=fine_giorno)
        c = ws.cell(row=TS_HEADER_GIORNI_ROW, column=base, value=g)
        c.font, c.alignment, c.fill = grassetto, centro, grigio
        for slot in griglia["slot_info"]:
            cols = celle_slot(griglia, g, slot)
            if len(cols) > 1:
                ws.merge_cells(start_row=TS_HEADER_SLOT_ROW, start_column=cols[0],
                               end_row=TS_HEADER_SLOT_ROW, end_column=cols[-1])
            c = ws.cell(row=TS_HEADER_SLOT_ROW, column=cols[0], value=slot["label"])
            c.font, c.alignment = Font(bold=True, size=9), centro
            for col in cols:
                ws.column_dimensions[get_column_letter(col)].width = 4.5

    for col_lettera, width in [(COL_TS_MESE, 12), (COL_TS_PROGETTO, 22),
                               (COL_TS_MANSIONE, 34), (COL_TS_TOTALE, 10)]:
        ws.column_dimensions[col_lettera].width = width

    # Data rows: same order as the timesheet (by month, then trainee)
    righe_proposta.sort(key=lambda r: (r["anno"], r["mese"], r["tirocinante"]))
    riga_out = TS_DATA_START_ROW
    for rp in righe_proposta:
        ws.cell(row=riga_out, column=column_index_from_string(COL_TS_MESE),
                value=f"{rp['mese']:02d}/{rp['anno']}")
        ws.cell(row=riga_out, column=column_index_from_string(COL_TS_PROGETTO),
                value=rp["progetto"])
        ws.cell(row=riga_out, column=column_index_from_string(COL_TS_MANSIONE),
                value=rp["mansione"])
        tot = ws.cell(row=riga_out, column=column_index_from_string(COL_TS_TOTALE),
                      value=len(rp["celle"]) * UNITA_BASE)
        tot.font = grassetto
        for _g, col in rp["celle"]:
            c = ws.cell(row=riga_out, column=col, value=UNITA_BASE)
            c.fill, c.alignment, c.border = giallo, centro, bordo
        riga_out += 1

    ws.freeze_panes = ws.cell(row=TS_DATA_START_ROW,
                              column=column_index_from_string(COL_TS_TOTALE) + 1)
    wb.save(percorso)


# ============================================================
# MAIN
# ============================================================

def main():
    report = Report()

    print("Reading the timesheet...")
    wb_ts = load_workbook(TIMESHEET_PATH, data_only=True)
    ws_ts = wb_ts[TIMESHEET_SHEET]

    griglia = rileva_griglia(ws_ts)
    print(f"Grid detected: day 1 at column {get_column_letter(griglia['prima_col'])}, "
          f"{len(griglia['slot_info'])} slots/day, "
          f"{griglia['slot_info'][0]['larghezza']} half cells per slot, "
          f"{griglia['passo_giorno']} columns per day, {griglia['num_giorni']} days.")

    tutoraggi = estrai_tutoraggi(ws_ts, report)
    print(f"Found {len(tutoraggi)} trainees in the timesheet.\n")

    print("Opening the internships file...")
    wb_tir = load_workbook(FILE_TIROCINI_PATH)
    ws_tir = wb_tir[FILE_TIROCINI_SHEET] if FILE_TIROCINI_SHEET else wb_tir.active
    mappa_mesi = costruisci_mappa_mesi(ws_tir)

    col_inizio = column_index_from_string(COL_DATA_INIZIO)
    col_fine = column_index_from_string(COL_DATA_FINE)
    col_stato = column_index_from_string(COL_STATO)
    col_nome = column_index_from_string(COL_NOME_TIROCINANTE)

    # Reverse check: trainees in the internships file with no rows in the timesheet
    for riga in range(TIROCINI_DATA_START_ROW, ws_tir.max_row + 1):
        v = ws_tir.cell(row=riga, column=col_nome).value
        if v is None:
            continue
        nv = normalizza_testo(v)
        if nv in tutoraggi:
            continue
        if any(punteggio_similarita(nv, k) >= SOGLIA_FUZZY for k in tutoraggi):
            continue
        report.avviso(f"'{str(v).strip()}' è nel file tirocini ma non ha nessuna riga "
                      "'Tirocinio ...' nel timesheet: controllare Mansione/nome "
                      "(normale se il tirocinio non è ancora iniziato).")

    # The "Tirocinio ..." rows of the tutoring entries to redistribute do
    # NOT count as occupancy: those hours will be removed from the timesheet.
    righe_da_ignorare = set()
    for d in tutoraggi.values():
        righe_da_ignorare |= d.get("righe_ts", set())

    cache_occupazione = {}

    def occupazione(anno, mese):
        key = (anno, mese)
        if key not in cache_occupazione:
            cache_occupazione[key] = occupazione_mese(ws_ts, griglia, anno, mese,
                                                      righe_da_ignorare)
        return cache_occupazione[key]

    righe_proposta_xlsx = []

    for nome_norm in sorted(tutoraggi, key=lambda k: tutoraggi[k]["nome_originale"]):
        dati = tutoraggi[nome_norm]
        nome = dati["nome_originale"]

        mesi_grezzi = {}
        for r in dati["righe"]:
            mesi_grezzi[(r["anno"], r["mese"])] = mesi_grezzi.get((r["anno"], r["mese"]), 0.0) + r["ore"]
        ore_totali = sum(mesi_grezzi.values())
        if ore_totali <= 0:
            report.avviso(f"'{nome}': totale ore = 0 nel timesheet, nessuna proposta calcolata.")
            continue

        riga_tir, nome_tir, era_fuzzy, punteggio = trova_riga_tirocinante(ws_tir, nome_norm)
        if riga_tir is None:
            report.errore(f"'{nome}' non trovato nel file dei tirocini (nemmeno in modo "
                          "approssimato): controllare manualmente.")
            continue

        note = []
        if era_fuzzy:
            note.append(f"Abbinato in modo approssimato ({punteggio:.0%}) a '{nome_tir}' "
                        "nel file tirocini: verificare che sia la persona giusta.")

        stato_val = ws_tir.cell(row=riga_tir, column=col_stato).value
        stato_norm = normalizza_testo(stato_val)
        data_interruzione = parse_data_generica(stato_val) if "interrotto" in stato_norm else None
        if "interrotto" in stato_norm:
            if data_interruzione:
                note.append(f"Tirocinio INTERROTTO il {data_interruzione.strftime('%d/%m/%Y')}: "
                            "la distribuzione si ferma a quella data.")
            else:
                note.append("Tirocinio INTERROTTO (data non riconosciuta): verificare a mano.")

        data_inizio = parse_data_generica(ws_tir.cell(row=riga_tir, column=col_inizio).value)
        data_fine = parse_data_generica(ws_tir.cell(row=riga_tir, column=col_fine).value)
        if data_interruzione and (data_fine is None or data_interruzione < data_fine):
            data_fine = data_interruzione
        if data_inizio and data_fine and data_fine < data_inizio:
            note.append(f"Data fine ({data_fine.strftime('%d/%m/%Y')}) precedente alla data "
                        "inizio nel file tirocini: probabile refuso, data fine ignorata.")
            data_fine = None

        # --- Starting month: ALWAYS from the internship start. ---
        # BUG FIXED (v2): previously the start was taken from the month in
        # which the hours were MARKED (min(mesi_grezzi)). If that month fell
        # towards the END of the internship, the start->end window was
        # extremely narrow and all the hours got compressed there (excess)
        # instead of being spread over the months. Now it always starts from
        # the actual internship start.
        anno_seg, mese_seg = min(mesi_grezzi)
        if data_inizio:
            anno_p, mese_p = data_inizio.year, data_inizio.month
            if (anno_seg, mese_seg) != (anno_p, mese_p):
                note.append(f"Le ore erano tutte concentrate a {mese_seg:02d}/{anno_seg}: "
                            f"la distribuzione parte dall'inizio del tirocinio "
                            f"({data_inizio.strftime('%d/%m/%Y')}) e si spalma su tutti i mesi.")
        else:
            anno_p, mese_p = anno_seg, mese_seg
            note.append("Data di inizio non disponibile: la distribuzione parte dal mese "
                        f"in cui le ore risultano segnate ({mese_seg:02d}/{anno_seg}).")

        idx_partenza = next((i for i, m in enumerate(mappa_mesi)
                             if (m["anno"], m["mese"]) == (anno_p, mese_p)), None)
        if idx_partenza is None:
            report.errore(f"'{nome}': mese di partenza {mese_p:02d}/{anno_p} fuori dalle "
                          "colonne del file tirocini.")
            continue

        idx_limite = len(mappa_mesi) - 1
        if RISPETTA_DATA_FINE and data_fine:
            idx_calc = next((i for i, m in enumerate(mappa_mesi)
                             if (m["anno"], m["mese"]) == (data_fine.year, data_fine.month)), None)
            if idx_calc is not None:
                idx_limite = idx_calc

        mesi_finestra = mappa_mesi[idx_partenza:idx_limite + 1]

        # Day bounds (start/end month) + callback that places the hours
        # while checking the REAL monthly occupancy (occupazione_mese).
        def giorno_bounds(anno, mese):
            gmin, gmax = 1, calendar.monthrange(anno, mese)[1]
            if data_inizio and (anno, mese) == (data_inizio.year, data_inizio.month):
                gmin = data_inizio.day
            if data_fine and (anno, mese) == (data_fine.year, data_fine.month):
                gmax = min(gmax, data_fine.day)
            return gmin, gmax

        def capacita_mese(anno, mese, tetto):
            gmin, gmax = giorno_bounds(anno, mese)
            occ = occupazione(anno, mese)   # set in cache (mutated by proponi_slot_mese)
            return proponi_slot_mese(occ, griglia, anno, mese, tetto, gmin, gmax)

        allocazioni, ore_non_distr, eccedenza = distribuisci_mesi(
            ore_totali, mesi_finestra, capacita_mese)
        if eccedenza > 0:
            note.append(f"I mesi disponibili non bastano con il tetto di {MAX_ORE_MESE}h/mese: "
                        f"{eccedenza:g}h extra sono state ripartite sugli stessi mesi "
                        "(alcuni mesi superano il tetto). Verificare che vada bene.")
        if ore_non_distr > 0:
            note.append(f"ATTENZIONE: {ore_non_distr:g}h non distribuibili: slot liberi "
                        "insufficienti su tutti i mesi del tirocinio.")

        righe_mese = []
        for alloc in allocazioni:              # proposals already prepared (no double call)
            a, m_, ore_a = alloc["anno"], alloc["mese"], alloc["ore"]
            proposte = alloc["proposte"]
            dettaglio = descrivi_proposte(proposte)
            righe_mese.append({"mese": f"{m_:02d}/{a}",
                               "ore": f"{ore_a:g} h",
                               "dettaglio": dettaglio})
            if proposte:
                righe_proposta_xlsx.append({
                    "anno": a, "mese": m_, "tirocinante": nome,
                    "progetto": dati["progetto"], "mansione": dati["mansione"],
                    "celle": [(g, col) for g, _lab, col in proposte],
                })

        report.sezioni.append({
            "nome": nome,
            "intestazione": (f"Totale {ore_totali:g} ore, attualmente segnate a "
                             f"{mese_seg:02d}/{anno_seg}, da distribuire su "
                             f"{len(allocazioni)} mesi."),
            "note": note,
            "righe_mese": righe_mese,
        })

    # --- File 1: Excel proposal with layout identical to the timesheet ---
    scrivi_proposta_xlsx(OUTPUT_PROPOSTA_XLSX, griglia, ws_ts, righe_proposta_xlsx)
    print(f"Proposal file (timesheet layout) saved to: {OUTPUT_PROPOSTA_XLSX}")

    # --- File 2: Word report (or txt if python-docx is missing) ---
    oggi = date.today()
    if scrivi_report_docx(report, OUTPUT_REPORT_DOCX, oggi):
        print(f"Word report saved to: {OUTPUT_REPORT_DOCX}")
    else:
        percorso_txt = re.sub(r"\.docx$", ".txt", OUTPUT_REPORT_DOCX)
        scrivi_report_txt(report, percorso_txt, oggi)
        print("python-docx not installed (py -m pip install python-docx --user): "
              f"report saved in text format to {percorso_txt}")

    # --- (optional) grezzo file as in the old version ---
    if GENERA_FILE_GREZZO:
        for nome_norm, dati in tutoraggi.items():
            riga_tir, *_ = trova_riga_tirocinante(ws_tir, nome_norm)
            if riga_tir is None:
                continue
            for m in mappa_mesi:
                ws_tir.cell(row=riga_tir, column=m["col_idx"]).value = None
            mesi_grezzi = {}
            for r in dati["righe"]:
                mesi_grezzi[(r["anno"], r["mese"])] = mesi_grezzi.get((r["anno"], r["mese"]), 0.0) + r["ore"]
            for (a, m_), ore in mesi_grezzi.items():
                col = next((mm["col_idx"] for mm in mappa_mesi
                            if (mm["anno"], mm["mese"]) == (a, m_)), None)
                if col:
                    ws_tir.cell(row=riga_tir, column=col).value = ore
        wb_tir.save(OUTPUT_PATH_GREZZO)
        print(f"'Grezzo' file saved to: {OUTPUT_PATH_GREZZO}")

    print(f"\nDone: {len(report.sezioni)} trainees processed, "
          f"{len(report.avvisi)} warnings, {len(report.errori)} errors.")


if __name__ == "__main__":
    try:
        main()
    except PermissionError as e:
        # File locked by Excel/another app: a clear message instead of a
        # traceback, so from VS Code (F5) it is immediately clear what to do.
        nome_file = str(e).split("'")[-2] if "'" in str(e) else ""
        print("\n" + "=" * 70)
        print("ERROR: FILE LOCKED")
        print("=" * 70)
        if nome_file:
            print(f"Cannot open/write the file:\n  {nome_file}")
        print("\nMost likely cause: the file is open in Excel or in another")
        print("application holding an exclusive read/write lock on it.")
        print("\n=> Close the file in Excel and run the script again (F5).")
        print("=" * 70)
    except FileNotFoundError as e:
        nome_file = str(e).split("'")[-2] if "'" in str(e) else ""
        print("\n" + "=" * 70)
        print("ERROR: FILE NOT FOUND")
        print("=" * 70)
        if nome_file:
            print(f"The file does not exist at the specified path:\n  {nome_file}")
        print("\n=> Check the paths (TIMESHEET_PATH / FILE_TIROCINI_PATH)")
        print("   at the top of the script.")
        print("=" * 70)